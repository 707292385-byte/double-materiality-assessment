#!/usr/bin/env python3
"""步骤 4：读回收数据、清洗、计分。

支持 .ndjson（飞书 record-list 导出）、.csv（飞书/问卷星导出）、.xlsx（问卷星导出，需 openpyxl）。
输出：scores.json（供审核页）、议题得分明细.csv、回收与清洗说明。

用法：
  python3 dm_score.py --project <项目目录> \
      --impact <影响问卷导出> --finance <财务问卷导出> [--min-duration 30]
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import statistics as st
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from dm_lib import (  # noqa: E402
    conclusion,
    financial_score,
    impact_direction,
    impact_to_output,
    load_project,
    save_project,
)

IMPACT_NAMES = {"基本无影响": 1, "影响较少": 2, "中等影响": 3, "较大影响": 4, "影响显著": 5}
LIKELY_NAMES = {"很低": 1, "低": 2, "中等": 3, "高": 4, "很高": 5}


def read_rows(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".ndjson" or suffix == ".jsonl":
        rows = []
        for line in path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line:
                rows.append(json.loads(line))
        return rows
    if suffix == ".csv":
        with open(path, encoding="utf-8-sig", newline="") as fh:
            return list(csv.DictReader(fh))
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise SystemExit("读取 xlsx 需要 openpyxl；或用 Excel 另存为 CSV 后再跑")
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        head = [str(c).strip() if c is not None else "" for c in rows[0]]
        return [dict(zip(head, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]
    raise SystemExit(f"不支持的回收文件格式：{path.suffix}")


def cell(row: dict, col: str):
    v = row.get(col)
    if isinstance(v, list):
        return v[0] if v else None
    return v


def to_number(v, kind: str = "num"):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s:
        return None
    m = re.match(r"^-?\d+(\.\d+)?", s)
    if m:
        return float(m.group(0))
    if kind == "impact_name" and s in IMPACT_NAMES:
        return float(IMPACT_NAMES[s])
    if kind == "likely_name" and s in LIKELY_NAMES:
        return float(LIKELY_NAMES[s])
    return None


def map_columns(header: list[str]) -> dict:
    imp, deg, lik, group, duration = {}, {}, {}, None, None
    for col in header:
        c = str(col)
        # 先认题号：Q28 利益相关方沟通 这类题干里带关键词的列不能被当成分组列
        m = re.match(r"\s*(Q\d+)\s*[\.、:：]?\s*(.*)$", c)
        if m:
            code, rest = m.group(1), m.group(2)
            if "可能性" in rest:
                lik[code] = c
            elif "影响程度" in rest:
                deg[code] = c
            else:
                imp[code] = c
            continue
        if any(k in c for k in ("利益相关方", "填报人身份", "受访者身份", "群体")):
            group = c
            continue
        if any(k in c for k in ("时长", "耗时", "duration")):
            duration = c
            continue
    return {"impact": imp, "degree": deg, "likelihood": lik, "group": group, "duration": duration}


def clean(rows: list[dict], cols: dict, topics: list[dict], min_duration: float) -> tuple[list[dict], dict]:
    kept, dropped = [], {"缺答": 0, "直线作答": 0, "答题过快": 0}
    for row in rows:
        if cols["duration"] and min_duration:
            d = to_number(cell(row, cols["duration"]))
            if d is not None and d < min_duration:
                dropped["答题过快"] += 1
                continue
        impact_vals, deg_vals, lik_vals = [], [], []
        for t in topics:
            code = t["code"]
            if code in cols["impact"]:
                v = to_number(cell(row, cols["impact"][code]), "num")
                if v is not None:
                    impact_vals.append(v)
            if code in cols["degree"]:
                d = to_number(cell(row, cols["degree"][code]), "impact_name")
                p = to_number(cell(row, cols["likelihood"].get(code, "")), "likely_name")
                if d is not None and p is not None:
                    deg_vals.append(d)
                    lik_vals.append(p)
        if not impact_vals and not deg_vals:
            dropped["缺答"] += 1
            continue
        same_impact = len(set(impact_vals)) <= 1 and len(impact_vals) > 1
        same_fin = len(set(deg_vals)) <= 1 and len(set(lik_vals)) <= 1 and len(deg_vals) > 1
        if same_impact and same_fin:
            dropped["直线作答"] += 1
            continue
        row["__group"] = cell(row, cols["group"]) if cols["group"] else ""
        row["__impact"] = impact_vals
        row["__deg"] = deg_vals
        row["__lik"] = lik_vals
        kept.append(row)
    return kept, dropped


def summarise(values: list[float]) -> dict:
    if not values:
        return {"n": 0, "mean": None, "sd": None}
    return {
        "n": len(values),
        "mean": round(st.fmean(values), 2),
        "sd": round(st.pstdev(values), 2) if len(values) > 1 else 0.0,
    }


def main():
    ap = argparse.ArgumentParser(description="双重重要性计分")
    ap.add_argument("--project", required=True)
    ap.add_argument("--impact", help="影响重要性问卷导出文件")
    ap.add_argument("--finance", help="财务重要性问卷导出文件")
    ap.add_argument("--min-duration", type=float, default=30, help="答题时长下限（秒），0 表示不剔除")
    ap.add_argument("--out", help="scores.json 输出路径")
    args = ap.parse_args()

    path = Path(args.project)
    workdir = path if path.is_dir() else path.parent
    project = load_project(path)
    topics = project["topics"]
    scoring = project["scoring"]
    t_imp = float(project["thresholds"]["impact"])
    t_fin = float(project["thresholds"]["financial"])

    resp = project.get("responses", {})
    impact_path = Path(args.impact or resp.get("impact") or "")
    finance_path = Path(args.finance or resp.get("financial") or "")
    for label, p in (("影响", impact_path), ("财务", finance_path)):
        if not str(p) or not p.exists():
            raise SystemExit(f"{label}问卷回收文件不存在：{p}（先导出答卷，或用 --impact/--finance 指定）")

    impact_rows = read_rows(impact_path)
    finance_rows = read_rows(finance_path)
    cols_i = map_columns(list(impact_rows[0].keys()) if impact_rows else [])
    cols_f = map_columns(list(finance_rows[0].keys()) if finance_rows else [])

    missing = [t["code"] for t in topics if t["code"] not in cols_i["impact"]]
    if missing:
        print(f"⚠ 影响问卷缺这些题号列：{missing[:8]}{' …' if len(missing) > 8 else ''}，请核对题号是否从 Q01 起连续")
    missing_f = [t["code"] for t in topics if t["code"] not in cols_f["degree"] or t["code"] not in cols_f["likelihood"]]
    if missing_f:
        print(f"⚠ 财务问卷缺这些题号列：{missing_f[:8]}{' …' if len(missing_f) > 8 else ''}")

    kept_i, dropped_i = clean(impact_rows, cols_i, topics, args.min_duration)
    kept_f, dropped_f = clean(finance_rows, cols_f, topics, args.min_duration)

    out_topics = []
    for t in topics:
        code = t["code"]
        imp_by_topic = [to_number(cell(r, cols_i["impact"][code]), "num") for r in kept_i if code in cols_i["impact"]]
        imp_by_topic = [v for v in imp_by_topic if v is not None]

        deg_by_topic, lik_by_topic = [], []
        for r in kept_f:
            if code in cols_f["degree"]:
                d = to_number(cell(r, cols_f["degree"][code]), "impact_name")
                p = to_number(cell(r, cols_f["likelihood"].get(code, "")), "likely_name")
                if d is not None and p is not None:
                    deg_by_topic.append(d)
                    lik_by_topic.append(p)

        imp_stat = summarise(imp_by_topic)
        fin_scores = [financial_score(d, p, scoring) for d, p in zip(deg_by_topic, lik_by_topic)]
        fin_stat = summarise(fin_scores)
        imp_score = impact_to_output(imp_stat["mean"], scoring) if imp_stat["mean"] is not None else None
        fin_score = fin_stat["mean"]

        out_topics.append(
            {
                "code": code,
                "name": t["name"],
                "category": t["category"],
                "description": t["description"],
                "impact": {
                    "n": imp_stat["n"],
                    "rawMean": imp_stat["mean"],
                    "rawSd": imp_stat["sd"],
                    "score": round(imp_score, 2) if imp_score is not None else None,
                    "direction": impact_direction(imp_stat["mean"], scoring) if imp_stat["mean"] is not None else "",
                },
                "finance": {
                    "n": fin_stat["n"],
                    "degreeMean": round(st.fmean(deg_by_topic), 2) if deg_by_topic else None,
                    "likelihoodMean": round(st.fmean(lik_by_topic), 2) if lik_by_topic else None,
                    "score": round(fin_score, 2) if fin_score is not None else None,
                    "sd": fin_stat["sd"],
                },
                "conclusion": conclusion(imp_score or 0, fin_score or 0, t_imp, t_fin),
                "note": "",
            }
        )

    ranked = sorted([t for t in out_topics if t["impact"]["score"] is not None and t["finance"]["score"] is not None],
                    key=lambda x: x["impact"]["score"] + x["finance"]["score"], reverse=True)
    for i, t in enumerate(ranked, start=1):
        t["rank"] = i
    for t in out_topics:
        t.setdefault("rank", None)

    divergence = []
    if cols_i["group"]:
        groups = sorted({r.get("__group") or "未填写" for r in kept_i})
        for t in out_topics:
            code = t["code"]
            if code not in cols_i["impact"]:
                continue
            per = []
            for g in groups:
                vals = [
                    to_number(cell(r, cols_i["impact"][code]), "num")
                    for r in kept_i
                    if (r.get("__group") or "未填写") == g
                ]
                vals = [v for v in vals if v is not None]
                if len(vals) >= 2:
                    per.append(st.fmean(vals))
            if len(per) >= 2:
                divergence.append({"code": code, "name": t["name"], "sd": round(st.pstdev(per), 2), "groups": len(per)})
        divergence.sort(key=lambda x: x["sd"], reverse=True)

    imp_scores = [t["impact"]["score"] for t in out_topics if t["impact"]["score"] is not None]
    fin_scores_all = [t["finance"]["score"] for t in out_topics if t["finance"]["score"] is not None]
    suggested = {
        "impact": round(st.median(imp_scores), 1) if imp_scores else t_imp,
        "financial": round(st.median(fin_scores_all), 1) if fin_scores_all else t_fin,
    }

    payload = {
        "schema": "dm-scores/1.0",
        "company": project["company"],
        "industry": project["industry"],
        "scoring": scoring,
        "thresholds": {"impact": t_imp, "financial": t_fin},
        "suggestedThresholds": suggested,
        "meta": {
            "impact": {"rows": len(impact_rows), "kept": len(kept_i), "dropped": dropped_i,
                       "source": str(impact_path), "groupColumn": cols_i["group"]},
            "finance": {"rows": len(finance_rows), "kept": len(kept_f), "dropped": dropped_f,
                        "source": str(finance_path)},
        },
        "topics": out_topics,
        "divergence": divergence[:8],
    }

    out = Path(args.out) if args.out else workdir / "scores.json"
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    detail = workdir / "议题得分明细.csv"
    with open(detail, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["排名", "编号", "议题", "类别", "影响样本数", "影响均值(原始)", "影响重要性得分(1-5)",
                    "财务样本数", "影响程度均值", "可能性均值", "财务重要性得分(1-5)", "结论"])
        for t in ranked:
            w.writerow([t.get("rank"), t["code"], t["name"], t["category"], t["impact"]["n"], t["impact"]["rawMean"],
                        t["impact"]["score"], t["finance"]["n"], t["finance"]["degreeMean"],
                        t["finance"]["likelihoodMean"], t["finance"]["score"], t["conclusion"]])

    project["responses"] = {"impact": str(impact_path), "financial": str(finance_path)}
    project["stage"] = "scored"
    save_project(project, workdir)

    dual = sum(1 for t in out_topics if t["conclusion"] == "双重重要")
    ratio = dual / max(1, len(out_topics))
    if ratio >= 0.6:
        print(f"⚠ 双重重要占比 {ratio:.0%} 偏高：1~5 原始均值下多数议题会同时越过 3.0 阈值。")
        print(f"  可把阈值调到中位数附近再审核：影响 {suggested['impact']} / 财务 {suggested['financial']}"
              f"（写入：dm_lib.py set --project <目录> thresholds.impact={suggested['impact']} thresholds.financial={suggested['financial']}）")
    print(f"影响问卷：{len(impact_rows)} 条 → 有效 {len(kept_i)}（剔除 {dropped_i}）")
    print(f"财务问卷：{len(finance_rows)} 条 → 有效 {len(kept_f)}（剔除 {dropped_f}）")
    print(f"议题 {len(out_topics)} 项｜双重重要 {dual} 项｜阈值 影响 {t_imp} / 财务 {t_fin}")
    print(f"输出：{out}\n      {detail}")


if __name__ == "__main__":
    main()
